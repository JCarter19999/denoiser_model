import argparse
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def _metrics(path: Path) -> dict:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)["overall"]


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    heading.style.font.color.rgb = RGBColor(0, 0, 0)


def _save_segmentation_plot(metrics: dict[str, dict], path: Path) -> None:
    labels = list(metrics)
    values = [metrics[label]["miou"] for label in labels]
    colors = ["#5B9BD5", "#ED7D31", "#70AD47"]
    figure, axis = plt.subplots(figsize=(8, 4.5))
    bars = axis.bar(labels, values, color=colors)
    axis.set_ylim(0, max(values) + 0.08)
    axis.set_ylabel("Mean Intersection over Union")
    axis.set_title("ACDC Validation Segmentation Performance")
    axis.grid(axis="y", alpha=0.25)
    for bar, value in zip(bars, values, strict=True):
        axis.text(bar.get_x() + bar.get_width() / 2, value + 0.01, f"{value:.4f}", ha="center")
    figure.tight_layout()
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _save_restoration_plot(metrics: dict[str, dict], path: Path) -> None:
    labels = list(metrics)
    figure, axes = plt.subplots(1, 3, figsize=(11, 4))
    fields = [("psnr", "PSNR ↑"), ("ssim", "SSIM ↑"), ("l1", "L1 ↓")]
    for axis, (field, title) in zip(axes, fields, strict=True):
        values = [metrics[label][field] for label in labels]
        bars = axis.bar(labels, values, color=["#A5A5A5", "#4472C4"])
        axis.set_title(title)
        axis.grid(axis="y", alpha=0.25)
        for bar, value in zip(bars, values, strict=True):
            axis.text(bar.get_x() + bar.get_width() / 2, value, f"{value:.3f}", ha="center", va="bottom")
    figure.suptitle("Synthetic Restoration Evaluation on ACDC Reference Images", y=1.03)
    figure.tight_layout()
    figure.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(figure)


def _save_pipeline_plot(path: Path) -> None:
    figure, axis = plt.subplots(figsize=(11, 2.6))
    axis.axis("off")
    stages = [
        ("ACDC adverse image", "#FCE4D6"),
        ("Residual\ndecorrupter", "#D9EAF7"),
        ("SegFormer\nsegmenter", "#E2F0D9"),
        ("Semantic mask\n+ metrics", "#FFF2CC"),
    ]
    for index, (label, color) in enumerate(stages):
        x = 0.03 + index * 0.245
        rectangle = plt.Rectangle((x, 0.32), 0.19, 0.35, facecolor=color, edgecolor="#4F81BD", linewidth=1.5)
        axis.add_patch(rectangle)
        axis.text(x + 0.095, 0.495, label, ha="center", va="center", fontsize=11, weight="bold")
        if index < len(stages) - 1:
            axis.annotate("", xy=(x + 0.24, 0.495), xytext=(x + 0.19, 0.495), arrowprops={"arrowstyle": "->", "lw": 2})
    axis.set_xlim(0, 1)
    axis.set_ylim(0, 1)
    figure.tight_layout()
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the project presentation report as a DOCX file.")
    parser.add_argument("--output", default="outputs/project_presentation_report.docx")
    parser.add_argument("--assets-dir", default="outputs/report_assets")
    args = parser.parse_args()
    output_path = Path(args.output)
    assets_dir = Path(args.assets_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    assets_dir.mkdir(parents=True, exist_ok=True)

    evaluation_dir = Path("outputs/presentation_evaluation")
    segmentation_metrics = {
        "Baseline": _metrics(evaluation_dir / "none.json"),
        "Restoration-only": _metrics(evaluation_dir / "learned_presentation_restoration_final.json"),
        "Task-aware": _metrics(evaluation_dir / "learned_presentation_task_aware.json"),
    }
    restoration_dir = Path("outputs/presentation_restoration_evaluation")
    restoration_metrics = {
        "1 epoch": json.loads((restoration_dir / "one_epoch_restoration.json").read_text()),
        "Multi-epoch": json.loads((restoration_dir / "best_restoration.json").read_text()),
    }
    segmentation_plot = assets_dir / "segmentation_comparison.png"
    restoration_plot = assets_dir / "restoration_comparison.png"
    pipeline_plot = assets_dir / "project_pipeline.png"
    _save_segmentation_plot(segmentation_metrics, segmentation_plot)
    _save_restoration_plot(restoration_metrics, restoration_plot)
    _save_pipeline_plot(pipeline_plot)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_heading("Task-Aware Image Decorruption", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Project overview, experimental goals, training results, and presentation artifacts")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    _add_heading(document, "Executive Summary")
    document.add_paragraph(
        "This project evaluates whether a lightweight learned restoration front end can improve semantic segmentation "
        "of adverse-condition driving scenes. The system was trained and evaluated on ACDC images across fog, night, "
        "rain, and snow conditions. The final presentation pipeline includes restoration training, SegFormer fine-tuning, "
        "task-aware fine-tuning, quantitative validation, and visual training artifacts."
    )

    _add_heading(document, "Goals and Aims")
    _add_bullet(document, "Establish a strong segmentation baseline on ACDC adverse-condition images.")
    _add_bullet(document, "Train a residual decoder that removes synthetic corruptions from clean reference images.")
    _add_bullet(document, "Fine-tune the decoder with semantic-task feedback to preserve segmentation-relevant content.")
    _add_bullet(document, "Compare baseline, reconstruction-only, and task-aware preprocessing on held-out validation data.")
    _add_bullet(document, "Provide transparent visual evidence through reconstruction, activation, and loss-landscape videos.")

    _add_heading(document, "Methodology")
    document.add_picture(str(pipeline_plot), width=Inches(6.7))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "The decorrupter predicts a bounded residual that is added to the corrupted image. Reconstruction training uses "
        "a combined L1, SSIM, and smoothness objective. Task-aware fine-tuning additionally uses frozen SegFormer logits "
        "to prioritize image details that matter for semantic segmentation."
    )

    _add_heading(document, "Training Setup")
    _add_bullet(document, "Data: ACDC train/validation splits; 1,600 labeled train images and 406 validation images.")
    _add_bullet(document, "Hardware: NVIDIA GeForce RTX 4070 Ti with mixed-precision training.")
    _add_bullet(document, "Restoration: 12 additional epochs beyond the seed checkpoint at 256×512 resolution.")
    _add_bullet(document, "Segmenter: 10 epochs; best validation mIoU = 0.5496.")
    _add_bullet(document, "Task-aware decorrupter: 8 epochs initialized from the multi-epoch restoration checkpoint.")

    _add_heading(document, "Segmentation Results")
    document.add_picture(str(segmentation_plot), width=Inches(6.4))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ["Method", "mIoU", "Pixel accuracy"], strict=True):
        cell.text = text
    for label, values in segmentation_metrics.items():
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{values['miou']:.4f}"
        cells[2].text = f"{values['pixel_accuracy']:.4f}"
    document.add_paragraph(
        "Interpretation: restoration-only preprocessing currently lowers mIoU. Task-aware fine-tuning largely recovers "
        "baseline segmentation performance (0.5494 versus 0.5496), showing that semantic feedback prevents most of the "
        "restoration-induced degradation."
    )

    _add_heading(document, "Restoration Results")
    document.add_picture(str(restoration_plot), width=Inches(6.6))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "The multi-epoch reconstruction model improved SSIM and L1 relative to the one-epoch seed model, while PSNR "
        "was slightly lower. This reflects the combined perceptual objective rather than a PSNR-only optimization target."
    )

    _add_heading(document, "Qualitative Evidence")
    reconstruction_image = Path("outputs/presentation_restoration_final/qualitative_severe_reconstruction.png")
    landscape_image = Path("outputs/presentation_restoration_final/post_training_loss_landscape.png")
    if reconstruction_image.is_file():
        document.add_picture(str(reconstruction_image), width=Inches(6.6))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph(
            "Severe composite synthetic corruption: darkening, fog, rain streaks, and sensor noise. "
            "The panel shows the trained model output against the clean reference."
        )
    if landscape_image.is_file():
        document.add_picture(str(landscape_image), width=Inches(6.6))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Post-training local loss landscape sampled around the trained restoration checkpoint.")

    _add_heading(document, "Presentation Artifacts")
    _add_bullet(document, "outputs/presentation_restoration_final/reconstruction_learning.mp4")
    _add_bullet(document, "outputs/presentation_restoration_final/loss_landscape_learning.mp4")
    _add_bullet(document, "outputs/presentation_restoration_final/activation_learning.mp4")
    _add_bullet(document, "outputs/presentation_evaluation/*.json")
    _add_bullet(document, "outputs/presentation_restoration_evaluation/*.json")

    _add_heading(document, "Conclusions and Next Steps")
    document.add_paragraph(
        "The project has progressed from a one-epoch demonstration to a reproducible multi-stage training pipeline with "
        "trained checkpoints, validation metrics, visual diagnostics, and presentation-ready artifacts. The main finding is "
        "that naïve reconstruction can harm segmentation, while task-aware optimization preserves baseline segmentation "
        "accuracy. The next experiment should tune the task-aware loss weights and use condition-stratified validation to "
        "seek a measurable improvement over the baseline, especially for night and snow scenes."
    )

    document.save(output_path)
    print(f"report={output_path}")
    print(f"assets={assets_dir}")


if __name__ == "__main__":
    main()
