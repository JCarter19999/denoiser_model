"""Build an editable deck that can be imported into Google Slides."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "task_aware_decorruption_results_slides.pptx"
LANDSCAPE = ROOT / "outputs" / "test_restoration" / "post_training_loss_landscape.png"

NAVY = RGBColor(17, 37, 65)
BLUE = RGBColor(33, 113, 181)
TEAL = RGBColor(31, 159, 148)
GOLD = RGBColor(238, 178, 17)
INK = RGBColor(35, 42, 52)
MUTED = RGBColor(91, 103, 116)
PALE = RGBColor(239, 246, 250)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text, x, y, w, h, *, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, *, size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = INK
        p.space_after = Pt(10)
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.65, 0.35, 11.9, 0.48, size=27, color=NAVY, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(0.94), Inches(1.35), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 0.65, 1.08, 11.8, 0.35, size=12, color=MUTED)


def add_footer(slide, number):
    add_text(slide, f"Task-Aware Image Decorruption  |  {number}", 0.65, 7.12, 12, 0.22, size=9, color=MUTED)


def add_card(slide, x, y, w, h, heading, body, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALE
    shape.line.color.rgb = RGBColor(210, 224, 234)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.10), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, heading, x + 0.28, y + 0.18, w - 0.45, 0.34, size=17, color=NAVY, bold=True)
    add_text(slide, body, x + 0.28, y + 0.66, w - 0.45, h - 0.80, size=13, color=INK)


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1: project overview
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = WHITE
    add_text(slide, "Task-Aware Image Decorruption", 0.7, 0.70, 8.9, 0.62, size=31, color=NAVY, bold=True)
    add_text(slide, "Can a learned restoration front end preserve the visual information that semantic segmentation needs in adverse driving conditions?", 0.72, 1.50, 7.2, 0.82, size=20, color=INK)
    add_text(slide, "Project aim", 0.72, 2.65, 2.2, 0.32, size=15, color=TEAL, bold=True)
    add_bullets(slide, [
        "Improve road-scene segmentation under rain, fog, snow, and nighttime conditions.",
        "Compare raw input, classical enhancement, reconstruction-only restoration, and task-aware restoration.",
        "Optimize not only for a visually clean image, but for downstream segmentation utility."
    ], 0.76, 3.08, 7.1, 2.15, size=17)
    add_card(slide, 8.55, 1.06, 3.95, 1.42, "Input", "Adverse-condition driving image", BLUE)
    add_card(slide, 8.55, 2.76, 3.95, 1.42, "Decorrupter", "Lightweight learned restoration front end", TEAL)
    add_card(slide, 8.55, 4.46, 3.95, 1.42, "Frozen segmenter", "Evaluates whether restoration helps the target task", GOLD)
    add_text(slide, "Primary outcome: held-out semantic segmentation mIoU, supported by PSNR/SSIM and efficiency measures.", 0.75, 6.28, 10.9, 0.42, size=15, color=MUTED)
    add_footer(slide, 1)

    # Slide 2: configuration
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Experiment configuration", "Current saved artifacts are one-epoch smoke tests: they validate the pipeline, not final comparative performance.")
    add_card(slide, 0.70, 1.65, 3.85, 2.18, "Data and corruption", "Dataset: ACDC driving imagery\nInput size: 256 x 512\nBatch size: 1\nSynthetic severity range: 0.25-0.80\nRandom seed: 42", BLUE)
    add_card(slide, 4.74, 1.65, 3.85, 2.18, "Restoration model", "Residual decorrupter\nBase channels: 24\nResidual scale: 0.25\n1 epoch; learning rate: 2e-4\nAMP + gradient clipping", TEAL)
    add_card(slide, 8.78, 1.65, 3.85, 2.18, "Restoration objective", "L1 reconstruction: 1.00\nSSIM: 0.25\nSmoothness: 0.01\nSmoke-test composite loss: 0.1363", GOLD)
    add_card(slide, 0.70, 4.28, 5.85, 1.55, "Segmentation backbone", "SegFormer-B0, initialized from Cityscapes-pretrained weights; 19 classes. One-epoch smoke test: training loss 0.3749 and validation mIoU 0.4942.", BLUE)
    add_card(slide, 6.74, 4.28, 5.89, 1.55, "Planned evaluation", "Hold the segmenter fixed while training restoration variants. Evaluate raw, classical, reconstruction-only, task-aware, and regularized task-aware methods on the same held-out set.", TEAL)
    add_text(slide, "Interpretation guardrail: no task-aware checkpoint or shared-condition comparison is available yet, so improvement claims are premature.", 0.78, 6.20, 11.7, 0.42, size=15, color=RGBColor(158, 76, 0), bold=True)
    add_footer(slide, 2)

    # Slide 3: actual loss landscape
    slide = prs.slides.add_slide(blank)
    add_title(slide, "What the local loss landscape shows", "Post-training reconstruction-loss slice around the restoration checkpoint")
    slide.shapes.add_picture(str(LANDSCAPE), Inches(0.58), Inches(1.48), width=Inches(8.1), height=Inches(3.55))
    add_card(slide, 8.96, 1.48, 3.70, 1.35, "How it is computed", "The model is perturbed along two normalized random parameter directions. At each grid point, the reconstruction objective is evaluated on a fixed synthetic batch.", BLUE)
    add_card(slide, 8.96, 3.05, 3.70, 1.35, "How to read it", "Purple/blue regions are lower loss; green/yellow regions are higher loss. The blue marker is the sampled global minimum; the red marker is the sampled global maximum.", TEAL)
    add_card(slide, 8.96, 4.62, 3.70, 1.35, "What this plot suggests", "The sampled neighborhood contains a broad low-loss basin, with sharper increases toward the edges and one localized high-loss region. This is a local diagnostic, not a proof of global optimality or generalization.", GOLD)
    add_text(slide, "Observed sampled loss range: approximately 0.058 to 0.096. Directions 1 and 2 are abstract parameter-space directions, not image features or input axes.", 0.70, 5.43, 7.95, 0.56, size=13, color=MUTED)
    add_footer(slide, 3)

    # Slide 4: presentable takeaway
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Takeaway and next experiment", "The infrastructure is working; the scientific comparison is the next milestone.")
    add_text(slide, "What has been demonstrated", 0.82, 1.55, 4.8, 0.35, size=19, color=NAVY, bold=True)
    add_bullets(slide, [
        "End-to-end segmenter and reconstruction training, checkpointing, and visualization are operational.",
        "The local loss surface provides a repeatable view of optimization behavior around a trained checkpoint.",
        "Recorded reconstruction, activation, and loss-landscape animations are available for qualitative reporting."
    ], 0.82, 2.02, 5.95, 3.0, size=17)
    add_text(slide, "What is still required", 7.05, 1.55, 4.8, 0.35, size=19, color=NAVY, bold=True)
    add_bullets(slide, [
        "Train the reconstruction-only and task-aware variants to convergence.",
        "Run every method on the same held-out adverse-condition images.",
        "Report mIoU overall/by class/by condition, PSNR/SSIM, and runtime tradeoffs."
    ], 7.05, 2.02, 5.45, 3.0, size=17)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(5.60), Inches(11.1), Inches(0.78))
    shape.fill.solid(); shape.fill.fore_color.rgb = NAVY; shape.line.fill.background()
    box = add_text(slide, "Decision criterion: task-aware decorruption is successful only if it improves held-out segmentation quality while retaining acceptable visual quality and latency.", 1.38, 5.80, 10.55, 0.36, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_footer(slide, 4)

    OUTPUT.parent.mkdir(exist_ok=True)
    prs.save(OUTPUT)
    print(f"saved={OUTPUT}")


if __name__ == "__main__":
    make_deck()
